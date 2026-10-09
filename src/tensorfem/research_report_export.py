"""Export one complete Markdown research report with embedded figures.

All formats use the same Markdown source. Missing images are an error.
Optional dependencies are installed with tensorfem[reports].
"""
from __future__ import annotations

import base64
import html
import os
import re
from pathlib import Path


def export_report(source, output_dir):
    import markdown
    from lxml import html as parser
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfbase.ttfonts import TTFont
    from matplotlib.font_manager import findfont
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                   TableStyle, Image, Preformatted)

    source, output_dir = Path(source), Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    body = markdown.markdown(source.read_text(), extensions=['tables', 'fenced_code'])
    tree = parser.fragment_fromstring(body, create_parent='div')
    # Resolve every figure before writing any export.
    for img in tree.iter('img'):
        path = (source.parent / img.get('src')).resolve()
        if not path.is_file():
            raise ValueError('missing report figure: '+str(path))
        img.set('data-local-path', str(path))
    document = Document()
    document.styles['Normal'].font.size = Pt(10)
    for style in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3'):
        font = document.styles[style].element.get_or_add_rPr()
        fonts = OxmlElement('w:rFonts'); fonts.set(qn('w:eastAsia'), 'SimSun')
        font.append(fonts)
    font_path = Path('/usr/share/fonts/truetype/droid/DroidSansFallbackFull.ttf')
    if font_path.is_file():
        pdfmetrics.registerFont(TTFont('ReportCJK', str(font_path)))
    else:
        pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
    font = 'ReportCJK' if font_path.is_file() else 'STSong-Light'
    # DroidSansFallback has Chinese glyphs but lacks Latin/digit glyphs.
    # Use a real Latin font for quantities and switch only CJK spans.
    pdfmetrics.registerFont(TTFont('ReportLatin', findfont('DejaVu Sans')))
    style = ParagraphStyle('body', fontName='ReportLatin', fontSize=9, leading=14,
                           wordWrap='CJK', spaceAfter=7, alignment=TA_LEFT)
    cell_style = ParagraphStyle('cell', parent=style, fontSize=7, leading=10, spaceAfter=0)
    story = []
    width = 500

    def text(element):
        return ''.join(element.itertext()).strip()

    def pdf_text(value):
        parts = re.split(r'([\u2e80-\u9fff\uf900-\ufaff\uff00-\uffef]+)', value)
        return ''.join(f'<font name="{font}">{html.escape(p)}</font>' if i % 2
                       else html.escape(p) for i, p in enumerate(parts)).replace('\n', '<br/>')

    def paragraph(value, heading=0):
        document.add_heading(value, level=min(heading, 3)) if heading else document.add_paragraph(value)
        pdfstyle = ParagraphStyle('heading', parent=style, fontSize=16-heading,
                                  leading=21-heading, spaceBefore=10) if heading else style
        story.append(Paragraph(pdf_text(value), pdfstyle))

    def emit(element):
        tag = element.tag
        if tag in ('h1', 'h2', 'h3', 'h4'):
            paragraph(text(element), int(tag[1]))
        elif tag == 'table':
            rows = [[text(c) for c in row] for row in element.iter('tr')]
            if not rows:
                return
            count = len(rows[0])
            table = document.add_table(rows=1, cols=count)
            table.style = 'Table Grid'
            for i, row in enumerate(rows):
                cells = table.rows[0].cells if i == 0 else table.add_row().cells
                for cell, value in zip(cells, row):
                    cell.text = value
            data = [[Paragraph(pdf_text(v), cell_style) for v in row] for row in rows]
            pdf_table = Table(data, colWidths=[width/count]*count, repeatRows=1)
            pdf_table.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), .4, colors.grey),
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e8edf3')),
                    ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                    ('TOPPADDING', (0, 0), (-1, -1), 5),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 5)]))
            story.extend([pdf_table, Spacer(1, 9)])
        elif tag == 'pre':
            value = text(element)
            document.add_paragraph(value, style='Normal')
            # Plain text, preserving code newlines while allowing long paths to wrap.
            story.append(Paragraph(pdf_text(value), cell_style))
        elif tag in ('ul', 'ol'):
            for i, child in enumerate(element):
                paragraph(('• ' if tag == 'ul' else f'{i+1}. ')+text(child))
        elif tag == 'blockquote':
            for child in element:
                emit(child)
        elif list(element.iter('img')):
            for img in element.iter('img'):
                path = img.get('data-local-path')
                document.add_picture(path, width=Inches(6.2))
                w, h = ImageReader(path).getSize()
                story.append(Image(path, width=width, height=width*h/w))
                if img.get('alt'):
                    paragraph(img.get('alt'))
        elif tag not in ('hr',):
            if text(element):
                paragraph(text(element))

    for element in tree:
        emit(element)
    docx_path = output_dir / (source.stem+'.docx')
    pdf_path = output_dir / (source.stem+'.pdf')
    html_path = output_dir / (source.stem+'.html')
    document.save(docx_path)
    SimpleDocTemplate(str(pdf_path), pagesize=(595.28, 841.89), leftMargin=47,
                      rightMargin=47, topMargin=40, bottomMargin=40).build(story)
    for img in tree.iter('img'):
        path = Path(img.attrib.pop('data-local-path'))
        img.set('src', 'data:image/png;base64,'+base64.b64encode(path.read_bytes()).decode())
    for anchor in tree.iter('a'):
        href = anchor.get('href', '')
        if href and not href.startswith(('#', 'http:', 'https:', 'mailto:')):
            target = (source.parent/href).resolve()
            anchor.set('href', os.path.relpath(target, output_dir.resolve()))
    title = text(tree.find('h1')) if tree.find('h1') is not None else source.stem
    html_path.write_text("<!doctype html><html lang='zh-CN'><meta charset='utf-8'>"
            '<title>'+html.escape(title)+"</title><style>body{font:16px sans-serif;"
            'max-width:1100px;margin:2em auto;line-height:1.6;padding:1em}'
            'table{border-collapse:collapse;width:100%;font-size:14px}'
            'td,th{border:1px solid #bbb;padding:.5em}img{max-width:100%}'
            'pre{white-space:pre-wrap;background:#f4f4f4;padding:1em}'
            '</style><body>'+parser.tostring(tree, encoding='unicode')+'</body></html>')
    return [html_path, pdf_path, docx_path]
